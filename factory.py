#!/usr/bin/env python3
"""
factory.py — AI-powered presentation builder using claude-opus-4-6 via Amazon Bedrock

Usage:
    python factory.py --topic "Zero Trust for Financial Services" --theme midnight
    python factory.py --from-pptx existing_deck.pptx --theme meridian
    python factory.py --from-json plan.json --theme teal
    python factory.py --topic "AWS Security Hub" --plan-only > plan.json
"""
import argparse
import json
import os
import sys
from pathlib import Path

import boto3
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ─────────────────────────────────────────────
# THEMES
# ─────────────────────────────────────────────

THEMES = {
    "meridian":  {"bg_dark": "0F1F3D", "bg_light": "F8F9FA", "bg_card": "FFFFFF", "bg_card_dark": "162040", "accent_1": "C0001E", "accent_2": "1D4ED8", "accent_3": "065F46", "text_dark": "FFFFFF", "text_light": "111827", "text_muted": "6B7280", "callout_bg": "0F1F3D", "callout_text": "E5E7EB", "font": "Calibri"},
    "midnight":  {"bg_dark": "1E2761", "bg_light": "F0F4FF", "bg_card": "FFFFFF", "bg_card_dark": "252D6E", "accent_1": "3B82F6", "accent_2": "8B5CF6", "accent_3": "06B6D4", "text_dark": "E0E7FF", "text_light": "1E2761", "text_muted": "6B7280", "callout_bg": "1E2761", "callout_text": "CADCFC", "font": "Calibri"},
    "teal":      {"bg_dark": "065A82", "bg_light": "F0FAFA", "bg_card": "FFFFFF", "bg_card_dark": "0A7AAF", "accent_1": "028090", "accent_2": "00A896", "accent_3": "02C39A", "text_dark": "E0F7FA", "text_light": "065A82", "text_muted": "64748B", "callout_bg": "065A82", "callout_text": "B2EBF2", "font": "Calibri"},
    "corporate": {"bg_dark": "212529", "bg_light": "F8F9FA", "bg_card": "FFFFFF", "bg_card_dark": "343A40", "accent_1": "495057", "accent_2": "0D6EFD", "accent_3": "198754", "text_dark": "F8F9FA", "text_light": "212529", "text_muted": "6C757D", "callout_bg": "212529", "callout_text": "CED4DA", "font": "Calibri"},
    "forest":    {"bg_dark": "1A3728", "bg_light": "F4FBF6", "bg_card": "FFFFFF", "bg_card_dark": "22472F", "accent_1": "2C5F2D", "accent_2": "97BC62", "accent_3": "F4A261", "text_dark": "E8F5E9", "text_light": "1A3728", "text_muted": "6B7280", "callout_bg": "1A3728", "callout_text": "C8E6C9", "font": "Calibri"},
}

# Set PPT_FACTORY_MODEL_ID to the Claude model or inference profile enabled in your Bedrock account.
BEDROCK_MODEL_ID = os.environ.get("PPT_FACTORY_MODEL_ID", "us.anthropic.claude-opus-4-6-20250514-v1:0")

W = Inches(10)
H = Inches(5.625)


def _log(msg: str) -> None:
    """Progress messages go to stderr so stdout stays clean for --plan-only JSON."""
    print(msg, file=sys.stderr)


# ─────────────────────────────────────────────
# SHAPE HELPERS
# ─────────────────────────────────────────────

def _rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

def _rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = _rgb(color); s.line.fill.background()
    return s

def _text(slide, x, y, w, h, text, font, pt, bold=False, color="000000", align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(pt)
    r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = _rgb(color)
    return tb

def _bg(slide, color):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = _rgb(color)

def _card(slide, x, y, w, h, t, accent):
    _rect(slide, x, y, w, h, t["bg_card"])
    _rect(slide, x, y, 0.06, h, accent)

def _badge(slide, x, y, w, h, text, t):
    b = slide.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = _rgb("F3F4F6")
    b.line.fill.background(); b.adjustments[0] = 0.5
    p = b.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = text
    r.font.name = t["font"]; r.font.size = Pt(8)
    r.font.bold = True; r.font.color.rgb = _rgb(t["text_light"])

def _accent(t, key):
    return {"accent_1": t["accent_1"], "accent_2": t["accent_2"], "accent_3": t["accent_3"]}.get(key, t["accent_1"])


# ─────────────────────────────────────────────
# SLIDE BUILDERS
# ─────────────────────────────────────────────

def _title_slide(prs, d, t):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, t["bg_dark"])
    _rect(slide, 0, 0, 0.18, 5.625, t["accent_1"])
    _rect(slide, 6.5, 0, 3.5, 0.12, t["accent_1"])
    _rect(slide, 0.45, 0.4, 0.9, 0.9, t["text_dark"])
    _text(slide, 0.45, 0.4, 0.9, 0.9, "P1", t["font"], 18, bold=True, color=t["bg_dark"], align=PP_ALIGN.CENTER)
    _text(slide, 0.45, 1.5, 5.0, 0.3, (d.get("subtitle") or "").upper(), t["font"], 9, bold=True, color=t["accent_1"])
    _text(slide, 0.45, 1.9, 9.1, 1.0, d.get("title", ""), t["font"], 40, bold=True, color=t["text_dark"])
    if d.get("subtitle"):
        _text(slide, 0.45, 2.95, 9.1, 0.5, d["subtitle"], t["font"], 16, color="A0AEC0")
    _rect(slide, 0.45, 3.55, 5.0, 0.03, "2D4A7A")
    _text(slide, 0.45, 3.75, 9.1, 0.32, d.get("speaker_line", "[Name]  ·  [Role]"), t["font"], 12, color="A0AEC0")
    slide.notes_slide.notes_text_frame.text = d.get("speaker_notes", "")


def _three_cards_slide(prs, d, t):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, t["bg_light"])
    _rect(slide, 0, 0, 10, 0.2, t["accent_1"])
    _text(slide, 0.4, 0.32, 9.2, 0.56, d.get("title", ""), t["font"], 28, bold=True, color=t["text_light"])
    if d.get("subtitle"):
        _text(slide, 0.4, 0.9, 9.2, 0.36, d["subtitle"], t["font"], 12, color=t["text_muted"])
    cards = d.get("cards", [])
    n = max(len(cards), 1)
    cw = round(9.2 / n - 0.15 * (n - 1) / n, 3)
    for i, card in enumerate(cards):
        cx = 0.4 + i * (cw + 0.15)
        ac = _accent(t, card.get("phase_color", "accent_1"))
        _card(slide, cx, 1.38, cw, 3.22, t, ac)
        if card.get("badge"):
            _badge(slide, cx + 0.12, 1.5, min(len(card["badge"]) * 0.09 + 0.3, cw - 0.2), 0.26, card["badge"], t)
        _text(slide, cx + 0.14, 2.06, cw - 0.22, 0.65, card.get("heading", ""), t["font"], 13, bold=True, color=t["text_light"])
        _text(slide, cx + 0.14, 2.76, cw - 0.22, 1.65, card.get("body", ""), t["font"], 10, color=t["text_muted"])
    if d.get("callout"):
        _rect(slide, 0.4, 4.73, 9.2, 0.58, t["callout_bg"])
        _text(slide, 0.55, 4.73, 9.0, 0.58, d["callout"], t["font"], 11, color=t["callout_text"])
    slide.notes_slide.notes_text_frame.text = d.get("speaker_notes", "")


def _two_column_slide(prs, d, t):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, "FFFFFF")
    _rect(slide, 0, 0, 10, 0.2, t["accent_1"])
    _text(slide, 0.4, 0.32, 9.2, 0.5, d.get("title", ""), t["font"], 26, bold=True, color=t["text_light"])
    left = d.get("left_panel", {})
    right = d.get("right_panel", {})
    _card(slide, 0.4, 0.94, 4.4, 3.5, t, t["accent_1"])
    _text(slide, 0.56, 1.02, 4.1, 0.36, left.get("heading", ""), t["font"], 13, bold=True, color=t["text_light"])
    for j, item in enumerate(left.get("items", [])[:5]):
        iy = 1.52 + j * 0.56
        _rect(slide, 0.56, iy + 0.06, 0.08, 0.08, t["accent_1"])
        _text(slide, 0.72, iy, 3.9, 0.22, item.get("label", ""), t["font"], 11, bold=True, color=t["text_light"])
        _text(slide, 0.72, iy + 0.24, 3.9, 0.26, item.get("description", ""), t["font"], 10, color=t["text_muted"])
    _rect(slide, 5.05, 0.94, 4.55, 3.5, t["bg_card_dark"])
    _text(slide, 5.2, 1.02, 4.2, 0.36, right.get("heading", ""), t["font"], 13, bold=True, color=t["text_dark"])
    for j, item in enumerate(right.get("items", [])[:5]):
        iy = 1.52 + j * 0.56
        _rect(slide, 5.2, iy + 0.06, 0.08, 0.08, t["accent_2"])
        _text(slide, 5.36, iy, 4.1, 0.22, item.get("label", ""), t["font"], 11, bold=True, color=t["text_dark"])
        _text(slide, 5.36, iy + 0.24, 4.1, 0.26, item.get("description", ""), t["font"], 10, color="A0AEC0")
    if d.get("quote"):
        _rect(slide, 0.4, 4.6, 9.2, 0.72, "FFF0F2")
        _text(slide, 0.55, 4.6, 9.0, 0.72, f"\"{d['quote']}\"", t["font"], 12, italic=True, color=t["text_light"])
    slide.notes_slide.notes_text_frame.text = d.get("speaker_notes", "")


def _timeline_slide(prs, d, t):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, t["bg_light"])
    _rect(slide, 0, 0, 10, 0.2, t["accent_1"])
    _text(slide, 0.4, 0.32, 9.2, 0.5, d.get("title", ""), t["font"], 24, bold=True, color=t["text_light"])
    phases = d.get("phases", [])
    n = max(len(phases), 1)
    pw = 2.85; gap = 0.125
    sx = (10 - (n * pw + (n - 1) * gap)) / 2
    _rect(slide, sx, 1.92, n * pw + (n - 1) * gap, 0.06, "E5E7EB")
    for i, ph in enumerate(phases):
        px = sx + i * (pw + gap)
        ac = _accent(t, ph.get("phase_color", "accent_1"))
        _rect(slide, px + pw / 2 - 0.14, 1.82, 0.28, 0.28, ac)
        _card(slide, px, 2.1, pw, 2.85, t, ac)
        _badge(slide, px + 0.12, 2.18, 0.85, 0.26, ph.get("label", f"Phase {i+1}"), t)
        _text(slide, px + 1.05, 2.2, 1.7, 0.24, ph.get("timeframe", ""), t["font"], 8, color=t["text_muted"])
        _text(slide, px + 0.12, 2.52, pw - 0.2, 0.32, ph.get("heading", ""), t["font"], 14, bold=True, color=t["text_light"])
        for j, item in enumerate(ph.get("items", [])[:5]):
            iy = 2.92 + j * 0.38
            _rect(slide, px + 0.16, iy + 0.04, 0.09, 0.09, ac)
            _text(slide, px + 0.32, iy, pw - 0.44, 0.34, item, t["font"], 10, color=t["text_light"])
        if ph.get("outcome"):
            _rect(slide, px, 4.88, pw, 0.52, "FFF0F2")
            _text(slide, px + 0.1, 4.88, pw - 0.16, 0.52, ph["outcome"], t["font"], 9, color=t["text_light"])
    slide.notes_slide.notes_text_frame.text = d.get("speaker_notes", "")


def _closing_slide(prs, d, t):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, t["bg_dark"])
    _rect(slide, 0, 0, 10, 0.2, t["accent_1"])
    _rect(slide, 0, 5.42, 10, 0.2, t["accent_1"])
    _text(slide, 0.4, 0.28, 9.2, 0.38, "Wrap-Up", t["font"], 10, bold=True, color=t["accent_1"])
    _text(slide, 0.4, 0.72, 9.2, 0.72, d.get("title", ""), t["font"], 28, bold=True, color=t["text_dark"])
    for i, card in enumerate(d.get("cards", [])[:3]):
        cx = 0.4 + i * 3.07
        _rect(slide, cx, 1.58, 2.95, 2.1, t["bg_card_dark"])
        _text(slide, cx + 0.15, 1.68, 0.5, 0.5, card.get("icon", "●"), t["font"], 20, color=t["accent_1"], align=PP_ALIGN.CENTER)
        _text(slide, cx + 0.76, 1.72, 2.06, 0.36, card.get("heading", ""), t["font"], 13, bold=True, color=t["text_dark"])
        _text(slide, cx + 0.12, 2.18, 2.72, 1.28, card.get("body", ""), t["font"], 10, color="A0AEC0")
    _text(slide, 0.4, 3.72, 9.2, 0.36, d.get("framework_title", "Key Takeaways"), t["font"], 13, bold=True, color=t["text_dark"])
    for j, item in enumerate(d.get("framework_items", [])[:3]):
        iy = 4.18 + j * 0.36
        _rect(slide, 0.4, iy + 0.02, 0.04, 0.24, t["accent_1"])
        _text(slide, 0.55, iy, 9.0, 0.3, item, t["font"], 11, color=t["text_dark"])
    slide.notes_slide.notes_text_frame.text = d.get("speaker_notes", "")


BUILDERS = {
    "title": _title_slide,
    "three_cards": _three_cards_slide,
    "two_column": _two_column_slide,
    "timeline": _timeline_slide,
    "closing": _closing_slide,
}


# ─────────────────────────────────────────────
# PLAN VALIDATION
# ─────────────────────────────────────────────
# The limits mirror what the slide builders can render. Builders slice lists
# (items[:5], cards[:3]), so anything over a limit would be dropped silently;
# validation turns that into an error instead.

ACCENTS = {"accent_1", "accent_2", "accent_3"}
MAX_SLIDES = 30


class PlanValidationError(ValueError):
    """Raised when a plan does not match the schema the slide builders need."""

    def __init__(self, errors: list[str]):
        super().__init__("invalid plan:\n  - " + "\n  - ".join(errors))
        self.errors = errors


def _check_str(errors, obj, key, where, required=False):
    val = obj.get(key)
    if val is None:
        if required:
            errors.append(f"{where}: missing '{key}'")
    elif not isinstance(val, str) or not val.strip():
        errors.append(f"{where}: '{key}' must be a non-empty string")


def _check_list(errors, obj, key, where, lo, hi, required=False):
    val = obj.get(key)
    if val is None:
        if required:
            errors.append(f"{where}: missing '{key}'")
        return []
    if not isinstance(val, list):
        errors.append(f"{where}: '{key}' must be a list")
        return []
    if not lo <= len(val) <= hi:
        errors.append(f"{where}: '{key}' has {len(val)} items, expected {lo}-{hi}")
    return val


def _check_objects(errors, items, where, required_keys, optional_keys=()):
    for j, item in enumerate(items, 1):
        at = f"{where} item {j}"
        if not isinstance(item, dict):
            errors.append(f"{at}: must be a JSON object")
            continue
        for key in required_keys:
            _check_str(errors, item, key, at, required=True)
        for key in optional_keys:
            _check_str(errors, item, key, at)
        color = item.get("phase_color")
        if color is not None and color not in ACCENTS:
            errors.append(f"{at}: phase_color {color!r} must be one of {sorted(ACCENTS)}")


def _check_strings(errors, items, where):
    for j, item in enumerate(items, 1):
        if not isinstance(item, str) or not item.strip():
            errors.append(f"{where} item {j}: must be a non-empty string")


def validate_plan(plan) -> list[str]:
    """Return every schema problem in a plan. An empty list means it is safe to render."""
    if not isinstance(plan, dict):
        return ["plan must be a JSON object"]
    errors = []
    _check_str(errors, plan, "title", "plan", required=True)
    slides = _check_list(errors, plan, "slides", "plan", 1, MAX_SLIDES, required=True)
    for i, slide in enumerate(slides, 1):
        where = f"slide {i}"
        if not isinstance(slide, dict):
            errors.append(f"{where}: must be a JSON object")
            continue
        layout = slide.get("layout")
        if layout not in BUILDERS:
            errors.append(f"{where}: unknown layout {layout!r}, expected one of {sorted(BUILDERS)}")
            continue
        where = f"slide {i} ({layout})"
        _check_str(errors, slide, "title", where, required=True)
        for key in ("subtitle", "speaker_notes"):
            _check_str(errors, slide, key, where)
        if layout == "title":
            _check_str(errors, slide, "speaker_line", where)
        elif layout == "three_cards":
            cards = _check_list(errors, slide, "cards", where, 1, 4, required=True)
            _check_objects(errors, cards, f"{where} cards", ("heading", "body"), ("badge",))
            _check_str(errors, slide, "callout", where)
        elif layout == "two_column":
            for side in ("left_panel", "right_panel"):
                panel = slide.get(side)
                if not isinstance(panel, dict):
                    errors.append(f"{where}: '{side}' must be a JSON object")
                    continue
                _check_str(errors, panel, "heading", f"{where} {side}", required=True)
                items = _check_list(errors, panel, "items", f"{where} {side}", 1, 5, required=True)
                _check_objects(errors, items, f"{where} {side}", ("label",), ("description",))
            _check_str(errors, slide, "quote", where)
        elif layout == "timeline":
            phases = _check_list(errors, slide, "phases", where, 1, 3, required=True)
            _check_objects(errors, phases, f"{where} phases", ("heading",), ("label", "timeframe", "outcome"))
            for j, phase in enumerate(phases, 1):
                if isinstance(phase, dict):
                    steps = _check_list(errors, phase, "items", f"{where} phase {j}", 0, 5)
                    _check_strings(errors, steps, f"{where} phase {j}")
        elif layout == "closing":
            cards = _check_list(errors, slide, "cards", where, 0, 3)
            _check_objects(errors, cards, f"{where} cards", ("heading",), ("body", "icon"))
            _check_str(errors, slide, "framework_title", where)
            takeaways = _check_list(errors, slide, "framework_items", where, 0, 3)
            _check_strings(errors, takeaways, f"{where} framework_items")
    if slides and isinstance(slides[0], dict) and slides[0].get("layout") != "title":
        errors.append("slide 1: a deck must start with a 'title' slide")
    if slides and isinstance(slides[-1], dict) and slides[-1].get("layout") != "closing":
        errors.append(f"slide {len(slides)}: a deck must end with a 'closing' slide")
    return errors


# ─────────────────────────────────────────────
# BEDROCK PROMPTS
# ─────────────────────────────────────────────

GENERATE_PROMPT = """You are a presentation architect. Return ONLY valid JSON — no markdown, no commentary.

Layout types: title | three_cards | two_column | timeline | closing
Always start with title, end with closing. Vary layouts — no two consecutive the same.
phase_color values: accent_1 | accent_2 | accent_3

Schema:
{
  "title": "deck title",
  "executive_summary": "one sentence",
  "slides": [
    {"number":1,"layout":"title","title":"...","subtitle":"...","speaker_line":"Name · Role","speaker_notes":"..."},
    {"number":2,"layout":"three_cards","title":"...","subtitle":"...","cards":[{"badge":"...","heading":"...","body":"2-3 sentences.","phase_color":"accent_1"}],"callout":"...","speaker_notes":"..."},
    {"number":3,"layout":"two_column","title":"...","left_panel":{"heading":"...","items":[{"label":"...","description":"..."}]},"right_panel":{"heading":"...","items":[{"label":"...","description":"..."}]},"quote":"...","speaker_notes":"..."},
    {"number":4,"layout":"timeline","title":"...","phases":[{"label":"Phase 1","timeframe":"0-90 Days","heading":"...","items":["..."],"outcome":"...","phase_color":"accent_1"}],"speaker_notes":"..."},
    {"number":5,"layout":"closing","title":"...","cards":[{"icon":"🎯","heading":"...","body":"..."}],"framework_title":"...","framework_items":["...","...","..."],"speaker_notes":"..."}
  ]
}"""

CONVERT_PROMPT = """You are a presentation architect. Reinterpret the provided slide content into the JSON schema below.
Preserve ALL original content. Map each source slide to the best-fit layout.
Return ONLY valid JSON — no markdown, no commentary.

Layout types: title | three_cards | two_column | timeline | closing
Always start with title, end with closing. Vary layouts — no two consecutive the same.
phase_color values: accent_1 | accent_2 | accent_3

Use the same schema as generation — same fields, same structure."""

# The conversion prompt reuses the generation schema so both paths produce the same shape.
CONVERT_PROMPT += "\n\n" + GENERATE_PROMPT[GENERATE_PROMPT.index("Schema:"):]


# ─────────────────────────────────────────────
# BEDROCK CALL
# ─────────────────────────────────────────────

def _call_bedrock(system: str, user_msg: str) -> dict:
    client = boto3.client(
        "bedrock-runtime",
        region_name=os.environ.get("AWS_DEFAULT_REGION", "us-east-1"),
    )
    response = client.invoke_model(
        modelId=BEDROCK_MODEL_ID,
        body=json.dumps({
            "anthropic_version": "bedrock-2023-05-31",
            "max_tokens": 4096,
            "system": system,
            "messages": [{"role": "user", "content": user_msg}],
        }),
        contentType="application/json",
        accept="application/json",
    )
    raw = json.loads(response["body"].read())["content"][0]["text"].strip()
    if raw.startswith("```"):
        raw = raw.split("\n", 1)[1].rsplit("```", 1)[0]
    return json.loads(raw)


def _plan_from_model(system: str, user_msg: str) -> dict:
    """Ask Bedrock for a plan and validate it, giving the model one chance to repair a bad response."""
    try:
        plan = _call_bedrock(system, user_msg)
        errors = validate_plan(plan)
    except json.JSONDecodeError as e:
        errors = [f"response was not valid JSON: {e}"]
    if not errors:
        return plan
    _log(f"[ppt-factory] Plan failed validation ({len(errors)} issue(s)) — asking the model to repair it...")
    repair_msg = (
        f"{user_msg}\n\nYour previous response was rejected for these reasons:\n- "
        + "\n- ".join(errors)
        + "\nReturn the complete corrected JSON only."
    )
    try:
        plan = _call_bedrock(system, repair_msg)
    except json.JSONDecodeError as e:
        raise PlanValidationError([f"response was not valid JSON: {e}"]) from e
    errors = validate_plan(plan)
    if errors:
        raise PlanValidationError(errors)
    return plan


# ─────────────────────────────────────────────
# PUBLIC FUNCTIONS
# ─────────────────────────────────────────────

def generate_plan(topic: str, audience: str = "technical", slide_count: int = 8) -> dict:
    """Generate a new deck plan from a topic via Bedrock."""
    _log(f"[ppt-factory] Calling Bedrock — {slide_count} slides for '{topic}'...")
    plan = _plan_from_model(
        GENERATE_PROMPT,
        f"Create a {slide_count}-slide presentation.\nTopic: {topic}\nAudience: {audience}\nReturn only JSON."
    )
    _log(f"[ppt-factory] Plan received: {len(plan['slides'])} slides")
    return plan


def extract_pptx_content(source_path: str) -> tuple[str, int]:
    """Extract text and speaker notes from an existing PPTX."""
    prs = Presentation(source_path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"\n--- SLIDE {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[NOTES]: {notes[:300]}")
    return "\n".join(lines), len(prs.slides)


def convert_plan(source_path: str) -> dict:
    """Extract content from an existing PPTX and reinterpret via Bedrock."""
    _log(f"[ppt-factory] Extracting: {source_path}")
    content, slide_count = extract_pptx_content(source_path)
    _log(f"[ppt-factory] Extracted {slide_count} slides — calling Bedrock...")
    plan = _plan_from_model(
        CONVERT_PROMPT,
        f"Source had {slide_count} slides. Reinterpret into structured JSON.\n\n{content}"
    )
    _log(f"[ppt-factory] Conversion plan: {len(plan['slides'])} slides")
    return plan


def build_pptx(plan: dict, theme_name: str = "meridian", output_path: str = "output.pptx") -> None:
    """Validate a plan, then render it to a .pptx file using the named theme."""
    if theme_name not in THEMES:
        raise ValueError(f"unknown theme {theme_name!r}, expected one of {sorted(THEMES)}")
    errors = validate_plan(plan)
    if errors:
        raise PlanValidationError(errors)
    t = THEMES[theme_name]
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    for i, slide_data in enumerate(plan["slides"], 1):
        layout = slide_data["layout"]
        BUILDERS[layout](prs, slide_data, t)
        _log(f"  [slide {i}] {layout} — {slide_data['title'][:55]}")
    prs.save(output_path)
    _log(f"[ppt-factory] Saved: {output_path}")


# ─────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="ppt-factory — AI presentation builder via Amazon Bedrock",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
examples:
  python factory.py --topic "Zero Trust for Financial Services" --theme midnight
  python factory.py --from-pptx existing.pptx --theme meridian
  python factory.py --topic "AWS Security Hub" --plan-only > plan.json
  python factory.py --from-json plan.json --theme teal
        """
    )
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--topic",     help="Generate new deck from topic")
    group.add_argument("--from-pptx", metavar="FILE", help="Convert existing .pptx to new theme")
    group.add_argument("--from-json", metavar="FILE", help="Build from existing plan JSON")

    parser.add_argument("--audience", default="technical",
                        choices=["technical", "executive", "mixed", "sales"])
    parser.add_argument("--slides",   type=int, default=8)
    parser.add_argument("--theme",    default="meridian", choices=list(THEMES.keys()))
    parser.add_argument("--output",   default=None)
    parser.add_argument("--plan-only", action="store_true",
                        help="Print JSON plan only, skip building .pptx")

    args = parser.parse_args()

    # Resolve output path
    if args.output:
        output_path = args.output
    elif args.from_pptx:
        output_path = f"{Path(args.from_pptx).stem}_{args.theme}.pptx"
    else:
        slug = (args.topic or "deck")[:40].lower().replace(" ", "_").replace("/", "_")
        output_path = f"{slug}_{args.theme}.pptx"

    # Get or load plan; every path is validated before anything is written
    try:
        if args.from_json:
            with open(args.from_json) as f:
                plan = json.load(f)
            _log(f"[ppt-factory] Loaded: {args.from_json}")
            errors = validate_plan(plan)
            if errors:
                raise PlanValidationError(errors)
        elif args.from_pptx:
            plan = convert_plan(args.from_pptx)
        else:
            plan = generate_plan(args.topic, args.audience, args.slides)
    except PlanValidationError as e:
        print(f"[ppt-factory] {e}", file=sys.stderr)
        sys.exit(2)

    if args.plan_only:
        print(json.dumps(plan, indent=2))
        sys.exit(0)

    # Save plan JSON
    plan_path = output_path.replace(".pptx", "_plan.json")
    with open(plan_path, "w") as f:
        json.dump(plan, f, indent=2)
    _log(f"[ppt-factory] Plan saved: {plan_path}")

    build_pptx(plan, args.theme, output_path)
